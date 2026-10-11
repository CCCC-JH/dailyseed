# -*- coding: utf-8 -*-
"""把 DailySeed 设计文档从 Markdown 转成排版良好的 DOCX。

只处理本文件实际用到的 Markdown 子集：
  标题 # ## ### ####  /  表格  /  围栏代码块  /  引用 >  /  列表 - 1.  /  分隔线 ---
  行内：**粗体**  `代码`  [文字](链接)
"""
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

CN_FONT = "Microsoft YaHei"
MONO_FONT = "Courier New"


def style_run(run, bold=False, mono=False, size=None, color=None, italic=False):
    name = MONO_FONT if mono else CN_FONT
    run.font.name = name
    run.font.bold = bold
    run.font.italic = italic
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name)


def shade_paragraph(paragraph, fill="F4F4F4"):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def shade_cell(cell, fill="EDEDED"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_inline(par, text):
    """把带行内标记的文本写进段落，保留 **粗体** / `代码` / [链接] 。"""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            style_run(par.add_run(part[2:-2]), bold=True)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            style_run(par.add_run(part[1:-1]), mono=True,
                      size=9, color=RGBColor(0xC0, 0x39, 0x2B))
        elif part.startswith("[") and "](" in part:
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            if m:
                style_run(par.add_run(m.group(1)), color=RGBColor(0x1A, 0x5F, 0xB4))
                style_run(par.add_run(" (" + m.group(2) + ")"),
                          size=7.5, color=RGBColor(0x88, 0x88, 0x88))
            else:
                style_run(par.add_run(part))
        else:
            style_run(par.add_run(part))


def add_code_block(doc, code_lines):
    for raw in code_lines:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.left_indent = Inches(0.18)
        pf.line_spacing = 1.0
        style_run(p.add_run(raw if raw.strip() else " "), mono=True, size=9)
        shade_paragraph(p)


def add_table(doc, rows):
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci in range(ncols):
            text = row[ci] if ci < len(row) else ""
            cell = cells[ci]
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_before = Pt(1)
            par.paragraph_format.space_after = Pt(1)
            add_inline(par, text)
            for run in par.runs:
                if run.font.size is None:
                    run.font.size = Pt(8.5)
                if ri == 0:
                    run.font.bold = True
            if ri == 0:
                shade_cell(cell)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def is_table_sep(line):
    s = line.strip()
    if not s.startswith("|"):
        return False
    return bool(re.fullmatch(r"\|[\s:\-\|]+\|", s))


def split_row(line):
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def convert(md_path, docx_path):
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().replace("\r\n", "\n").split("\n")

    doc = Document()

    # 全局默认字体（含中文 eastAsia）
    normal = doc.styles["Normal"]
    normal.font.name = CN_FONT
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for sec in doc.sections:
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        # 围栏代码块
        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            add_code_block(doc, buf)
            continue

        # 表格
        if stripped.startswith("|") and i + 1 < n and is_table_sep(lines[i + 1]):
            rows = [split_row(line)]
            i += 2
            while i < n and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(10)
                style_run(p.add_run(text), bold=True, size=20,
                          color=RGBColor(0x14, 0x3D, 0x2E))
            else:
                style = {2: "Heading 1", 3: "Heading 2", 4: "Heading 3"}[level]
                h = doc.add_paragraph(style=style)
                for r in list(h.runs):
                    r.text = ""
                h.paragraph_format.space_before = Pt(14 if level == 2 else 10)
                h.paragraph_format.space_after = Pt(6)
                sizes = {2: 15, 3: 12.5, 4: 11}
                colors = {2: RGBColor(0x14, 0x3D, 0x2E),
                          3: RGBColor(0x1F, 0x5C, 0x45),
                          4: RGBColor(0x33, 0x33, 0x33)}
                style_run(h.add_run(text), bold=True, size=sizes[level],
                          color=colors[level])
            i += 1
            continue

        # 分隔线
        if re.fullmatch(r"-{3,}", stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "CCCCCC")
            pbdr.append(bottom)
            pPr.append(pbdr)
            i += 1
            continue

        # 引用块（可多行）
        if stripped.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(re.sub(r"^\s*>\s?", "", lines[i]))
                i += 1
            for b in buf:
                if not b.strip():
                    continue
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.space_after = Pt(3)
                inner = b.strip()
                mm = re.match(r"^(#{2,4})\s+(.*)$", inner)
                if mm:
                    style_run(p.add_run(mm.group(2)), bold=True, size=12)
                else:
                    add_inline(p, inner)
                for r in p.runs:
                    if r.font.color.rgb is None:
                        r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            continue

        # 无序列表
        m = re.match(r"^[-*]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, m.group(1))
            i += 1
            continue

        # 有序列表
        m = re.match(r"^\d+[.)]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, m.group(1))
            i += 1
            continue

        # 空行
        if not stripped:
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_inline(p, stripped)
        i += 1

    doc.save(docx_path)
    print("saved:", docx_path)


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
